# -*- coding: utf-8 -*-
"""清洗层：已验收原件 → `PreparedBatch`（规范化行 + 门店匹配 + 去重）。

对应 spec §5.1/§5.2 与实施计划 T4。

复用（**不重写**）
- `normalize`：`parse_cell` / `parse_number` / `round_money` / `is_summary_row` /
  `output_colname` / 多行合并感知表头解析 `_expand_merges` + `_split_header`
- `dedup`：`dedup_table`（短键 + 批次、评价逐列覆盖、售后只打标记）
- `sync_store_master`：门店 ID/别名索引与匹配
- `shangou_merge.read_csv_text`：闪购 CSV 编码识别

新增（本模块）
- 规则来自 `references/cleaning_rules.json`，**逐表声明**，不套统一规则
- 只消费**清单列明的原件**，不扫描旧目录拼批次（无 glob 残留）
- 未决业务含义的列**不猜值**，标 `pending`
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import dedup as D            # noqa: E402
import normalize as N        # noqa: E402
import sync_store_master as M  # noqa: E402
from common import contracts as C  # noqa: E402
from common import dates     # noqa: E402

DEFAULT_RULES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                             "references", "cleaning_rules.json")
STORE_MASTER_DEFAULT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                    "references", "store_master.example.json")

# 待核标记：列名出现这些关键词且规则未给出正式来源时，值保持**待核**
PENDING_MARK = "pending"

# 事实表名统一加 `f_` 前缀 —— 与历史库（`import_legacy` 建的 `f_xxx`）保持一致。
# 🔴 这里不能省：`db.apply_batch` 直接拿 `group["table"]` 建表/写入，若输出不带前缀的
#    表名，新数据会落进**另一套同名无前缀的表**，历史数据仍在 `f_` 表里 ——
#    两边割裂、查询口径不一致，且极难发现（2026-10-02 实机踩到）。
FACT_PREFIX = "f_"


def fact_table(name):
    """业务表名 → 库内事实表名（幂等：已有前缀不重复加）。"""
    return name if name.startswith(FACT_PREFIX) else FACT_PREFIX + name


def load_rules(path=None):
    rules = C.read_json(path or DEFAULT_RULES)
    if "tables" not in rules or not rules["tables"]:
        raise ValueError("清洗规则缺 tables")
    return rules


def load_master(path=None):
    """门店主档：`{"stores": [...]}`；缺文件时返回空档（**不猜 ID**）。"""
    p = path or STORE_MASTER_DEFAULT
    if not p or not os.path.exists(p):
        return {"stores": []}
    return C.read_json(p)


def master_indexes(master):
    stores = (master or {}).get("stores") or []
    return M.build_id_mapping(stores), M.build_name_index(stores), M.normalize_store_name


# ---------------- 读原件 ----------------
_FNAME_DATE = re.compile(r"_(\d{8})")


def _read_xlsx(path, header_rows):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    grid, spans, raw = N._expand_merges(ws, header_rows)
    hdr = N._split_header(grid, spans, raw)
    data = [list(r) for r in ws.iter_rows(min_row=header_rows + 1, values_only=True)]
    data = [r + [None] * (len(hdr) - len(r)) for r in data]
    # `sheetnames` 是 **workbook** 的属性，worksheet 上没有（此前误写 `ws.sheetnames`，
    # 会在读第一个 xlsx 时抛 AttributeError）。
    return hdr, [[N._clean(c) for c in r[:len(hdr)]] for r in data], wb.sheetnames


def _read_csv(path):
    from shangou_merge import read_csv_text
    hdr, rows = read_csv_text(path)
    return [N.normalize_colname(c) for c in hdr], rows, ["csv"]


def read_raw(path, table, rules):
    """读一批原件 → `(列名, 数据行, 工作表, 文件名日期)`。"""
    r = rules["tables"][table]
    base = os.path.basename(path)
    m = _FNAME_DATE.search(base)
    fname_date = dates.to_dashed(m.group(1)) if m else None
    if path.lower().endswith((".xlsx", ".xlsm")):
        hdr, rows, sheets = _read_xlsx(path, r["header_rows"])
    else:
        hdr, rows, sheets = _read_csv(path)
    return hdr, rows, sheets, fname_date


# ---------------- 清洗一批 ----------------
def prepare_group(task, result, rules, master_idx, *, collected_at=None):
    """一个任务的原件 → 一个 group（含清洗记录与来源定位）。"""
    table = task["table"]
    rule = rules["tables"][table]
    id_map, name_idx, norm_name = master_idx
    collected_at = collected_at or result.get("collected_at") or ""

    if result["status"] == C.S_NO_DATA:
        # 无事件也是**可替换的修订**：不能因为「没有记录」就省略
        return {"task_key": task["task_key"], "platform": task["platform"],
                "table": fact_table(table), "target_table": fact_table(table),
                "business_date": task["business_date"], "window_start": task.get("window_start"),
                "window_end": task.get("window_end"), "scope_id": task["scope_id"],
                "status": C.S_NO_DATA, "columns": [], "rows": [],
                "source_files": [], "source_rows": 0, "cleaning_log": [],
                "business_key": list(rule["business_key"]), "snapshot": rule["snapshot"],
                "load_to_db": rule["load_to_db"],
                "optional_keys": list(rule.get("optional_keys") or [])}

    drop = set(rule["drop_cols"])
    date_candidates = rule.get("date_source_candidates") or []
    all_rows, cols, log, sheets_seen, src_rows = [], None, [], set(), 0

    for f in result.get("files") or []:
        path = f["path"]
        hdr, rows, sheets, fname_date = read_raw(path, table, rules)
        sheets_seen.update(sheets)
        biz_cols = [c for c in hdr if c not in drop and c != N.BATCH_COL]
        if cols is None:
            out_cols = [N.output_colname(c) for c in biz_cols]
            have_biz_date = "业务日期" in out_cols
            src = [c for c in N.SRC_COLS if not (c == "业务日期" and have_biz_date)]
            cols = out_cols + src + [N.SUMMARY_COLS, "门店标准ID"]
        for raw in rows:
            if not any(N._clean(x) for x in raw):
                continue
            src_rows += 1
            row = dict(zip(hdr, raw))
            rec = {N.output_colname(c): N.parse_cell(row.get(c), c, table) for c in biz_cols}
            if rule["date_source"] == "filename":
                day = fname_date or ""
            else:
                day = next((N.to_date(row.get(c)) for c in date_candidates
                            if N.to_date(row.get(c))), "")
            rec["业务日期"] = day
            rec["数据平台"] = rule["platform"]
            rec["数据表"] = table
            rec["采集时间"] = N._clean(row.get(N.BATCH_COL)) or fname_date or collected_at
            rec["渠道口径"] = rule["channel_scope"]
            is_sum = N.is_summary_row(row, table)
            rec[N.SUMMARY_COLS] = 1 if is_sum else 0
            sid = N.match_store_row(row, id_map, name_idx, norm_name)
            rec["门店标准ID"] = sid or ""
            all_rows.append(rec)

    # 去重（逐表声明：短键+批次 / 评价覆盖 / 售后打标 / 其余透传）
    deduped, dropped, extra = D.dedup_table(table, all_rows) if all_rows else ([], [], {})
    # 售后表**始终**带标记列，避免「有无错位行」导致两批列不一致
    if table == "售后订单数据" and cols:
        if D.AFTER_SALE_FLAG not in cols:
            cols = cols + [D.AFTER_SALE_FLAG]
        for r in deduped:
            r.setdefault(D.AFTER_SALE_FLAG, "0")
        log.append({"dedup": "售后列段标记", "stats": extra})

    pending_cols = [c for c in (cols or []) if "未决" in c or "待核" in c]
    for c in pending_cols:
        log.append({"pending_column": c, "reason": "业务含义未决，不自选正式来源"})

    log.append({"dedup": {"input": src_rows, "kept": len(deduped), "dropped": len(dropped)},
                "sheets": sorted(sheets_seen)})

    return {
        "task_key": task["task_key"], "platform": task["platform"],
        "table": fact_table(table), "target_table": fact_table(table),
        "business_date": task["business_date"],
        "window_start": task.get("window_start"), "window_end": task.get("window_end"),
        "scope_id": task["scope_id"], "status": C.S_VALIDATED,
        "columns": cols or [], "rows": [[r.get(c, "") for c in (cols or [])] for r in deduped],
        "source_files": [{"path": f["path"], "sha256": f.get("sha256"),
                          "bytes": f.get("bytes")} for f in (result.get("files") or [])],
        "source_rows": src_rows, "dedup_dropped": len(dropped),
        "cleaning_log": log, "business_key": list(rule["business_key"]),
        "snapshot": rule["snapshot"], "load_to_db": rule["load_to_db"],
        # 🔴 必透传：规则把「门店标准ID」声明为 optional_keys（汇总行门店id=-999 匹配不到店，
        # 属**正常**），`db.validate_group` 靠它区分「允许为空」与「键缺失」。
        # 之前这里漏传 → 汇总行全被判 EMPTY_KEY_VALUE，整批拒绝入库。
        "optional_keys": list(rule.get("optional_keys") or []),
    }


def prepare_batch(manifest, rules, master, *, hist=None):
    """BatchManifest → PreparedBatch。**只消费清单列明的原件**，不扫目录。"""
    C.validate_manifest(manifest)
    if not C.is_collection_complete(manifest):
        raise ValueError("本批次未全部合格，不得进入清洗：%s" % manifest["collection_status"])

    by_key = {t["task_key"]: t for t in manifest["tasks"]}
    m_idx = master_indexes(master)
    groups, sources, missing = [], [], []
    for r in manifest["results"]:
        task = by_key[r["task_key"]]
        for f in r.get("files") or []:
            if not os.path.exists(f["path"]):
                missing.append(f["path"])
    # 原件在清洗前消失 → **先拒绝，再动手**（不留下半成品）
    if missing:
        raise ValueError("原件在清洗前消失，拒绝继续：%s" % missing)

    for r in manifest["results"]:
        task = by_key[r["task_key"]]
        groups.append(prepare_group(task, r, rules, m_idx))
        sources.extend({"path": f["path"], "sha256": f.get("sha256"),
                        "table": task["table"], "status": r["status"]}
                       for f in (r.get("files") or []))

    prepared = {
        "batch_id": manifest["batch_id"],
        "business_today": manifest["business_today"],
        "rule_version": rules["rule_version"],
        "source_files": sources,
        "groups": groups,
        "stats": {"groups": len(groups),
                  "load_to_db": sum(1 for g in groups if g["load_to_db"]),
                  "no_data": sum(1 for g in groups if g["status"] == C.S_NO_DATA),
                  "rows": sum(len(g["rows"]) for g in groups)},
    }
    return prepared


def process_batch(manifest, rules, master, *, db_path=None, t=None, hist=None):
    """一次调用完成「清洗 + 入库」。

    - 无 `db_path`：只返回 `PreparedBatch`（独立清洗，不写库）
    - 有 `db_path`：调用 `db.apply_batch`，返回 `WriteResult`
    """
    prepared = prepare_batch(manifest, rules, master, hist=hist)
    if not db_path:
        return prepared
    import db as DB
    return DB.apply_batch(db_path, prepared, t or manifest["business_today"])
