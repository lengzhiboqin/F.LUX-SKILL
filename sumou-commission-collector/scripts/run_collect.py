#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""素某财务后台 · 第三方佣金明细采集技能（sumou-commission-collector）

一键闭环：连接CDP浏览器 → 打开素某后台财务中心-代扣打款 → 设置时间范围（上月1号~本月1号）
→ 导出表格 → Excel加工（佣金列右侧新增"实到金额"列，值=佣金-0.5）
→ 重命名"<上月>月第三方佣金明细.xlsx" → 上传企业微信微盘「豆包的共享空间」根目录。

登录：cookie持久化登录态优先；失效时自动读取 ~/.config/sumou-collector/credentials.json
（权限600）以手机号+密码自动登录；仍失败则退出码2请求人工接管。

用法：
  python3 run_collect.py                     # 自动按今天计算上月1号~本月1号
  python3 run_collect.py --date 20261001      # 指定执行日期（计算其上月1号~本月1号）
  python3 run_collect.py --out <目录>         # 指定本地输出目录（默认 项目目录/产出物/第三方佣金）

退出码：
  0 = 全流程成功
  1 = 页面/导出/加工/上传失败（详见摘要）
  2 = 未登录且自动登录失败（需人工接管浏览器登录后重跑）
  3 = 环境错误（缺依赖/CDP不可用）
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright

CDP_URL = "http://127.0.0.1:9222"
BASE_URL = "https://console.sumou.fun/#/finance/flow"
LOGIN_URL = "https://console.sumou.fun/#/login?redirect=%2Ffinance%2Fflow"
SPACE_KEYWORD = "豆包"
DEFAULT_OUT = str(Path.home() / "Doubao/chats/38439034639855618/产出物/第三方佣金")
CRED_FILE = str(Path.home() / ".config/sumou-collector/credentials.json")


def log(msg):
    print(msg, flush=True)


def load_credentials():
    """安全读取本地凭据（权限600），返回 dict 或 None。凭据不输出、不入日志。"""
    try:
        st = os.stat(CRED_FILE)
        if st.st_mode & 0o077:  # 权限过宽则视为不安全
            os.chmod(CRED_FILE, 0o600)
        with open(CRED_FILE, encoding="utf-8") as f:
            d = json.load(f)
        u, p = (d.get("username") or "").strip(), d.get("password") or ""
        return {"username": u, "password": p} if u and p else None
    except Exception:
        return None


def auto_login(page):
    """自动登录素某后台（手机号+密码+勾选协议+立即登录）。成功返回True。"""
    cred = load_credentials()
    if not cred:
        return False
    try:
        page.wait_for_timeout(2000)
        phone = page.locator("input[type=number][placeholder*='手机号']").first
        pwd = page.locator("input[type=password][placeholder*='密码']").first
        phone.fill(cred["username"])
        pwd.fill(cred["password"])
        # 勾选"已阅读并同意用户协议"checkbox（Element UI隐藏原生input，需点可见label容器）
        cb_label = page.locator("label.el-checkbox").first
        if cb_label.count():
            orig = page.locator("input[type=checkbox]").first
            if orig.count() and not orig.is_checked():
                cb_label.click()
                page.wait_for_timeout(1000)
        page.wait_for_timeout(500)
        btn = page.locator("button", has_text="立即登录").first
        btn.click()
        page.wait_for_timeout(8000)
        return "login" not in page.url
    except Exception:
        return False


def run_wecom(args_list):
    """调用 wecom-cli，返回 (returncode, stdout)"""
    r = subprocess.run(args_list, capture_output=True, text=True, timeout=120)
    return r.returncode, r.stdout


def find_wecom_space():
    """搜索微盘共享空间，返回 space_id（内部ID，不外露）"""
    rc, out = run_wecom(["wecom-cli", "disk", "files", "search",
                         "--json", json.dumps({"keywords": [SPACE_KEYWORD],
                                               "search_type": "space", "limit": 5})])
    if rc != 0:
        raise RuntimeError(f"微盘搜索失败: {out[:300]}")
    data = json.loads(out)
    for f in data.get("files", []):
        if f.get("type") == "space" and SPACE_KEYWORD in (f.get("file_name") or ""):
            return f["id"]
    raise RuntimeError(f"未找到微盘共享空间「{SPACE_KEYWORD}」")


def upload_wecom(space_id, file_path, file_name):
    """上传文件到微盘空间根目录，返回 (ok, msg)"""
    rc, out = run_wecom(["wecom-cli", "disk", "files", "upload", "--json",
                         json.dumps({"folder_id": space_id, "file_name": file_name,
                                     "file_path": file_path})])
    if rc != 0:
        return False, out[:300]
    try:
        data = json.loads(out)
        f = data.get("file", {})
        return True, f"{f.get('file_name')} ({f.get('file_size')}B) -> {f.get('path')}"
    except Exception:
        return True, out[:200]


def process_excel(raw_xls_path, out_path):
    """加工Excel：佣金列右侧新增'实到金额'列（值=佣金-0.5）。返回 (rows, header)"""
    # 导出的 .xls 实际是 xlsx zip 内容，先复制成 .xlsx 再读
    tmp_xlsx = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    shutil.copy(raw_xls_path, tmp_xlsx)
    try:
        wb = openpyxl.load_workbook(tmp_xlsx)
        ws = wb.active
        header = [c.value for c in ws[1]]
        comm_col = None
        for i, h in enumerate(header, 1):
            if h == "佣金":
                comm_col = i
                break
        if comm_col is None:
            raise RuntimeError(f"未找到佣金列，表头: {header}")
        new_col = comm_col + 1
        ws.insert_cols(new_col)
        ws.cell(row=1, column=new_col, value="实到金额")
        total_val = 0.0
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=comm_col).value
            if v is None:
                continue
            try:
                nv = round(float(v) - 0.5, 4)
                ws.cell(row=r, column=new_col, value=nv)
                total_val += nv
            except Exception:
                ws.cell(row=r, column=new_col, value=None)
        # 实到金额列最后一行下方写入总计行（红色+26号字体）
        total_row = ws.max_row + 1
        tcell = ws.cell(row=total_row, column=new_col,
                        value=f"总计金额：{round(total_val, 4)}")
        tcell.font = Font(size=26, color="FF0000", bold=True)
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        wb.save(out_path)
        rows = ws.max_row - 1
        return rows, header
    finally:
        os.unlink(tmp_xlsx)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", help="执行日期 YYYYMMDD（默认今天）")
    ap.add_argument("--out", default=DEFAULT_OUT, help="本地输出目录")
    args = ap.parse_args()

    today = datetime.strptime(args.date, "%Y%m%d").date() if args.date else date.today()
    # 上月1号 ~ 本月1号
    first_this = today.replace(day=1)
    prev_month_end = first_this - timedelta(days=1)
    first_prev = prev_month_end.replace(day=1)
    start_str = first_prev.strftime("%Y-%m-%d")
    end_str = first_this.strftime("%Y-%m-%d")
    month_label = f"{first_prev.month}月"  # 文件名取上月月份（数据反映上月佣金）
    file_name = f"{month_label}第三方佣金明细.xlsx"

    result = {"ok": False, "date": today.strftime("%Y-%m-%d"),
              "start": start_str, "end": end_str, "month": month_label,
              "file_name": file_name}

    # 1. 连接CDP浏览器
    try:
        pw = sync_playwright().start()
        browser = pw.chromium.connect_over_cdp(CDP_URL)
    except Exception as e:
        log(json.dumps({"ok": False, "error": f"CDP连接失败: {str(e)[:200]}"}, ensure_ascii=False))
        sys.exit(3)

    try:
        ctx = browser.contexts[0] if browser.contexts else browser.new_context()
        page = ctx.pages[0] if ctx.pages else ctx.new_page()

        # 2. 打开代扣打款页
        page.goto(BASE_URL, wait_until="domcontentloaded", timeout=45000)
        page.wait_for_timeout(4000)

        # 3. 登录态检查：未登录则自动登录（手机号+密码）
        if "login" in page.url:
            if auto_login(page):
                page.goto(BASE_URL, wait_until="domcontentloaded", timeout=45000)
                page.wait_for_timeout(4000)
            else:
                result["error"] = "未登录且自动登录失败（缺少凭据或登录被拦截）"
                log(json.dumps(result, ensure_ascii=False))
                sys.exit(2)
        if "login" in page.url:
            result["error"] = "自动登录后仍未进入后台"
            log(json.dumps(result, ensure_ascii=False))
            sys.exit(2)

        # 4. 设置日期范围（antd DatePicker：fill + Enter）
        start_input = page.locator("input[placeholder*='开始日期']").first
        end_input = page.locator("input[placeholder*='结束日期']").first
        start_input.click()
        page.wait_for_timeout(500)
        start_input.fill(start_str)
        start_input.press("Enter")
        page.wait_for_timeout(1500)
        end_input.click()
        page.wait_for_timeout(500)
        end_input.fill(end_str)
        end_input.press("Enter")
        page.wait_for_timeout(2500)

        # 5. 点击导出并等待下载
        with page.expect_download(timeout=90000) as dl_info:
            page.locator("button", has_text="导出").click()
        dl = dl_info.value
        raw_path = dl.path()
        page.wait_for_timeout(2000)

        # 6. Excel加工
        out_path = os.path.join(args.out, file_name)
        rows, header = process_excel(raw_path, out_path)
        result["rows"] = rows
        result["excel_path"] = out_path
        result["header"] = header

        # 7. 上传企微微盘（只上传最终加工表）
        space_id = find_wecom_space()
        ok, msg = upload_wecom(space_id, out_path, file_name)
        result["wecom_upload"] = {"ok": ok, "msg": msg}
        if not ok:
            result["error"] = f"微盘上传失败: {msg}"
            log(json.dumps(result, ensure_ascii=False))
            sys.exit(1)

        # 8. 上传成功后立即删除本地表（不留本地残留，避免重复上传/混淆）
        try:
            if os.path.exists(out_path):
                os.unlink(out_path)
                result["local_cleaned"] = True
        except Exception as e:
            result["local_cleaned"] = False
            result["local_clean_error"] = str(e)[:120]

        result["ok"] = True
        log(json.dumps(result, ensure_ascii=False, indent=2))
        sys.exit(0)
    finally:
        pw.stop()


if __name__ == "__main__":
    main()
