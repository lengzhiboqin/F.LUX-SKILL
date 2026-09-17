#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""CSV总表 → Excel多sheet导出。

将 by_table/ 下的8张CSV总表合并为一个Excel文件，每张表一个sheet，
保留"更新日志"sheet。输出UTF-8编码的.xlsx文件。

用法：
  python3 export_excel.py --master <总表目录> --output <输出.xlsx>
  python3 export_excel.py --master <总表目录> --output <输出.xlsx> --log "本次更新说明"
  python3 export_excel.py --config <整合技能配置文件>  # 从配置读取路径
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime
from pathlib import Path

# 8张标准表
TABLES = [
    "门店财务明细",
    "门店成交明细",
    "商品数据",
    "问题订单数据",
    "流量明细(新)",
    "流量渠道明细(新)",
    "评价数据",
    "售后订单数据",
]

LOG_SHEET = "更新日志"


def read_csv(path):
    """读取CSV，自动判别编码(UTF-8-BOM/GBK)"""
    raw = Path(path).read_bytes()
    # 尝试UTF-8-BOM
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("gbk")
        except UnicodeDecodeError:
            text = raw.decode("utf-8", errors="replace")
    rows = list(csv.reader(text.splitlines()))
    return [r for r in rows if any(c.strip() for c in r)]


def detect_number(val):
    """尝试将字符串转为数字，失败返回原字符串"""
    if val is None or val == "":
        return ""
    # 整数
    try:
        if str(val).lstrip("-").isdigit():
            return int(val)
    except Exception:
        pass
    # 浮点数
    try:
        f = float(val)
        if f == int(f) and abs(f) < 1e15:
            return int(f)
        return f
    except (ValueError, TypeError):
        pass
    return val


def export_to_excel(master_dir, output_path, log_message=None):
    """导出Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("错误：需要 openpyxl 库，运行: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    master_dir = Path(master_dir)
    by_table = master_dir / "by_table"
    output_path = Path(output_path)

    if not by_table.exists():
        print(f"错误：总表目录不存在: {by_table}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.Workbook()
    # 删除默认sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    exported = []
    for table_name in TABLES:
        csv_path = by_table / f"{table_name}.csv"
        if not csv_path.exists():
            print(f"  ⚠ 跳过(文件不存在): {table_name}")
            continue

        rows = read_csv(csv_path)
        if len(rows) < 1:
            print(f"  ⚠ 跳过(空文件): {table_name}")
            continue

        # 创建sheet（sheet名最长31字符）
        sheet_name = table_name[:31]
        ws = wb.create_sheet(title=sheet_name)

        header = rows[0]
        body = rows[1:]

        # 写表头
        for col_idx, val in enumerate(header, 1):
            cell = ws.cell(row=1, column=col_idx, value=val)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # 写数据（尝试数字转换）
        for row_idx, row in enumerate(body, 2):
            for col_idx, val in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=detect_number(val))

        # 冻结首行
        ws.freeze_panes = "A2"

        # 自动列宽（简单估算）
        for col_idx in range(1, len(header) + 1):
            max_len = len(str(header[col_idx - 1])) if col_idx <= len(header) else 8
            # 抽样前100行估算
            for row in body[:100]:
                if col_idx <= len(row):
                    cell_len = len(str(row[col_idx - 1]))
                    if cell_len > max_len:
                        max_len = cell_len
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 30)

        exported.append((table_name, len(body)))
        print(f"  ✓ {table_name}: {len(body)}行")

    # 更新日志sheet
    log_ws = wb.create_sheet(title=LOG_SHEET, index=0)
    log_ws.cell(row=1, column=1, value="日期").font = header_font
    log_ws.cell(row=1, column=1).fill = header_fill
    log_ws.cell(row=1, column=2, value="更新说明").font = header_font
    log_ws.cell(row=1, column=2).fill = header_fill
    log_ws.freeze_panes = "A2"

    # 读取已有日志（如果存在）
    existing_logs = []
    log_csv = by_table / "_update_log.csv"
    if log_csv.exists():
        existing_logs = read_csv(log_csv)
        if existing_logs and existing_logs[0] == ["日期", "更新说明"]:
            existing_logs = existing_logs[1:]

    # 写入已有日志
    for row_idx, log_row in enumerate(existing_logs, 2):
        log_ws.cell(row=row_idx, column=1, value=log_row[0] if len(log_row) > 0 else "")
        log_ws.cell(row=row_idx, column=2, value=log_row[1] if len(log_row) > 1 else "")

    # 追加本次日志
    if log_message:
        next_row = len(existing_logs) + 2
        log_ws.cell(row=next_row, column=1, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        log_ws.cell(row=next_row, column=2, value=log_message)
        # 同步保存到CSV
        with open(log_csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["日期", "更新说明"])
            for log_row in existing_logs:
                w.writerow(log_row)
            w.writerow([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), log_message])

    log_ws.column_dimensions["A"].width = 20
    log_ws.column_dimensions["B"].width = 80

    # 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))

    print(f"\n✓ Excel导出完成: {output_path}")
    print(f"  共 {len(exported)} 张表，{sum(r for _, r in exported)} 行数据")
    return output_path, exported


def main():
    parser = argparse.ArgumentParser(description="CSV总表 → Excel多sheet导出")
    parser.add_argument("--master", help="总表目录(含by_table子目录)")
    parser.add_argument("--output", help="输出Excel文件路径")
    parser.add_argument("--config", help="整合技能配置文件路径(从中读取master_dir)")
    parser.add_argument("--log", help="本次更新说明(写入更新日志)")
    args = parser.parse_args()

    # 从配置读取
    if args.config and Path(args.config).exists():
        cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
        if not args.master:
            args.master = cfg.get("master_dir")
        if not args.output and args.master:
            args.output = str(Path(args.master) / "F.LUX品牌数据汇总表.xlsx")

    if not args.master:
        # 尝试默认配置（新管线配置优先，回退旧版）
        for cfg_name in ["shangou-pipeline", "shangou-integrator"]:
            default_config = Path.home() / ".config" / cfg_name / "config.json"
            if default_config.exists():
                cfg = json.loads(default_config.read_text(encoding="utf-8"))
                args.master = cfg.get("master_dir")
                args.output = str(Path(args.master) / "F.LUX品牌数据汇总表.xlsx")
                break

    if not args.master:
        print("错误：请指定 --master 或 --config", file=sys.stderr)
        sys.exit(1)

    if not args.output:
        args.output = str(Path(args.master) / "F.LUX品牌数据汇总表.xlsx")

    print(f"总表目录: {args.master}")
    print(f"输出文件: {args.output}")
    print()

    export_to_excel(args.master, args.output, args.log)


if __name__ == "__main__":
    main()
